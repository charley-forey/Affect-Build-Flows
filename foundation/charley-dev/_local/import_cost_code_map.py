"""Regenerate the old -> new cost-code seeds from the client's mapping workbook.

    python import_cost_code_map.py "<path>/Cost Code Mapping_OldvsNew_v1.xlsx"
    python import_cost_code_map.py --legacy "<path>/Legacy Procore Codes to New_v1_CE.xlsx"
    python import_cost_code_map.py --master "<path>/COST CODE_WBS MASTER.xlsx"
    python make_qc_seeds.py            # then inline the CSVs as 09_seed_costcode.sql

The first writes 02-transformation/seed/cost_code_new.csv and cost_code_map.csv; --legacy
writes cost_code_legacy_map.csv (either or both in one run); --master rewrites only
cost_code_new.csv from the client's master list, which supersedes 'New Cost Codes'. The client will send
v2/v3 of the workbook; re-running this and make_qc_seeds.py is the whole update, and the
diff of the two CSVs is the review.

WHICH COLUMNS, AND WHY NOT THE OTHERS.
  'Cost Code Mapping'  Old Code, Old Description, New Code (header row 4). The Old columns
                       are formulas into 'OLD CC', so the workbook is read data_only - the
                       cached values, as the client last saw them.
  'New Cost Codes'     code (A), description (B), division (D, "1 - GENERAL REQUIREMENTS"),
                       data from row 3.
The mapping sheet's 'New Description' column is ignored: in v1 it is literal 'Not uploaded'
text, not a lookup. Descriptions and divisions come only from 'New Cost Codes'.

A blank New Code is written as a blank, not dropped: v1 leaves 340000.000 / 340001.000
(ALLOWANCES) without a target, and an old code missing from the seed would read as "not an
old code" rather than as "an old code nobody has mapped yet".
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
SEED_DIR = HERE.parent / "02-transformation" / "seed"


def code(value) -> str:
    """'10111.000' as text. A cell Excel re-typed as a number comes back as a float."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return f"{value:.3f}"
    return str(value).strip()


def read(xlsx: Path) -> tuple[list[dict], list[dict]]:
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    # Fail on a moved header rather than importing the wrong column silently.
    header = next(wb["Cost Code Mapping"].iter_rows(min_row=4, max_row=4, values_only=True))
    assert header[0] == "Old Code" and header[3] == "New Code", f"mapping header moved: {header}"

    new = []
    for row in wb["New Cost Codes"].iter_rows(min_row=3, values_only=True):
        if not code(row[0]):
            continue
        division = str(row[3] or "").strip()
        m = re.match(r"^(\d{1,2})\s*-\s*(.+)$", division)
        assert m, f"new code {row[0]}: division {division!r} is not 'N - NAME'"
        new.append({"new_cost_code": code(row[0]), "description": str(row[1] or "").strip(),
                    "division_code": m.group(1).zfill(2), "division_name": m.group(2).strip()})

    old = [{"old_cost_code": code(r[0]), "new_cost_code": code(r[3]),
            "old_description": str(r[1] or "").strip()}
           for r in wb["Cost Code Mapping"].iter_rows(min_row=5, values_only=True) if code(r[0])]
    return new, old


def read_master(xlsx: Path) -> list[dict]:
    """The new code list from 'Affect CC_WBS_UPLOAD' in the client's COST CODE_WBS MASTER.

    Only rows marked Uploaded to SAGE = YES: the Division 00 design codes are on the sheet but
    'No', and a code Sage does not have is not one Procore or Sage can post to. Data from row 3.
    """
    import openpyxl

    ws = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)["Affect CC_WBS_UPLOAD"]
    header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert header[0] == "Combined Code" and header[4] == "SAGE", f"master header moved: {header}"
    new = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not code(row[0]) or str(row[4] or "").strip().upper() != "YES":
            continue
        m = re.match(r"^(\d{1,2})\s*-\s*(.+)$", str(row[3] or "").strip())
        assert m, f"new code {row[0]}: division {row[3]!r} is not 'N - NAME'"
        new.append({"new_cost_code": code(row[0]), "description": str(row[1] or "").strip(),
                    "division_code": m.group(1).zfill(2), "division_name": m.group(2).strip()})
    return new


def read_legacy(xlsx: Path) -> list[dict]:
    """Pre-2026 Procore CSI codes -> new code, from the 'Legacy to New' sheet we sent the client.

    The legacy code is kept as the exact Procore string ('09-20-00', '01-00 00'): it is matched
    against Procore's code text, not normalised. Column G is the client's answer; column E was
    only our suggestion and is ignored.
    """
    import openpyxl

    ws = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)["Legacy to New"]
    header = next(ws.iter_rows(min_row=4, max_row=4, values_only=True))
    assert header[0] == "Legacy Code" and header[6] == "New Code", f"legacy header moved: {header}"
    return [{"legacy_cost_code": str(r[0]).strip(), "new_cost_code": code(r[6]),
             "legacy_description": str(r[1] or "").strip()}
            for r in ws.iter_rows(min_row=5, values_only=True) if r[0] is not None and str(r[0]).strip()]


def write(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path, nargs="?", help="old -> new mapping workbook")
    parser.add_argument("--legacy", type=Path, help="legacy Procore -> new mapping workbook")
    parser.add_argument("--master", type=Path, help="COST CODE_WBS MASTER workbook (new codes)")
    args = parser.parse_args()
    if not (args.xlsx or args.legacy or args.master):
        parser.error("give the mapping workbook, --legacy, --master, or a combination")
    if args.xlsx and args.master:
        parser.error("the mapping workbook and --master both write cost_code_new.csv; pick one")

    if args.master:
        new = read_master(args.master)
        write(SEED_DIR / "cost_code_new.csv", new)
        print(f"cost_code_new.csv  {len(new)} new codes (uploaded to Sage)")

    if args.legacy:
        legacy = read_legacy(args.legacy)
        write(SEED_DIR / "cost_code_legacy_map.csv", legacy)
        with (SEED_DIR / "cost_code_new.csv").open(encoding="utf-8") as fh:
            known = {r["new_cost_code"] for r in csv.DictReader(fh)}
        dangling = sorted({r["legacy_cost_code"] + " -> " + r["new_cost_code"] for r in legacy
                           if r["new_cost_code"] and r["new_cost_code"] not in known})
        print(f"cost_code_legacy_map.csv  {len(legacy)} legacy codes")
        if dangling:
            print(f"  WARNING legacy targets absent from New Cost Codes: {dangling}")
    if not args.xlsx:
        print("now run make_qc_seeds.py to regenerate 09_seed_costcode.sql")
        return 0

    new, old = read(args.xlsx)
    write(SEED_DIR / "cost_code_new.csv", new)
    write(SEED_DIR / "cost_code_map.csv", old)

    known = {r["new_cost_code"] for r in new}
    unmapped = [r["old_cost_code"] for r in old if not r["new_cost_code"]]
    dangling = sorted({r["new_cost_code"] for r in old} - known - {""})
    print(f"cost_code_new.csv  {len(new)} new codes")
    print(f"cost_code_map.csv  {len(old)} old codes, {len(old) - len(unmapped)} mapped")
    print(f"  unmapped old codes: {unmapped}")
    # Also a DQ warning in gold; printed here so it is seen before anything is committed.
    if dangling:
        print(f"  WARNING map targets absent from New Cost Codes: {dangling}")
    print("now run make_qc_seeds.py to regenerate 09_seed_costcode.sql")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
